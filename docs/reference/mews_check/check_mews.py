"""materials_mews_ashrae.xlsx の MEWS の材料（No.1～42）を、RR-110 の表の値と照らし合わせる。

実行方法（プロジェクトのフォルダで）:
    python docs/reference/mews_check/check_mews.py

結果は標準出力に表示する。docs/reference/mews_excel_derivation.md の数値は、この出力から取った。
"""

import math
import os
import statistics
import sys

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..', '..'))
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

from modules.material_props import PowerExpIsotherm, PowerExpVaporConductivity  # noqa: E402
import rr110_tables as T  # noqa: E402

EXCEL = os.path.join(ROOT, 'modules', 'materials_mews_ashrae.xlsx')

# 吸引圧 → 相対湿度（RR-110 の式(4)）に使う値。試験は室温 22 °C 前後。
_M = 0.018015    # 水のモル質量, kg/mol
_R = 8.314       # 気体定数, J/(mol·K)
_RHO_W = 1000.0  # 水の密度, kg/m3
_T_LAB = 295.15  # 22 °C, K


# ---------------------------------------------------------------------------
# Excel の読み込みと、Excel の式の計算
# ---------------------------------------------------------------------------

def read_excel() -> dict:
    """「データファイル用」シートを {No: 行の値の辞書} で返す。"""

    ws = openpyxl.load_workbook(EXCEL, data_only=True)['データファイル用']
    rows = {}
    for r in range(3, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 45)]
        if v[0] is None:
            continue
        rows[v[0]] = {
            'name': v[1], 'rho': v[2], 'c': v[3], 'lambda': v[4], 'k_air': v[5],
            'vapour': tuple(v[6:10]),
            'iso_parameter': v[10], 'iso': tuple(v[11:17]),
            'n_ranges': v[17],
            'ranges': [(None, v[18], tuple(v[19:26])), (v[26], v[27], tuple(v[28:35])), (v[35], v[36], tuple(v[37:44]))],
        }
    return rows


def vapour(row, x):
    a, b, c, d = row['vapour']
    return PowerExpVaporConductivity(a=a, b=b, c=c, d=d).lambda_dsh_m(x)


def isotherm(row, x):
    a, b, c, d, f, g = row['iso']
    if row['iso_parameter'] == 2:
        return PowerExpIsotherm(a=a, b=b, c=c, d=d, f=f, g=g).u(x)
    return PowerExpIsotherm(a=a, b=b, c=c, d=d).u(x)


def poly(coefficients, u):
    a, b, c, d, e, f, g = coefficients
    return a + b * u + c * u**2 + d * u**3 + e * u**4 + f * u**5 + g * u**6


def diffusivity(row, u):
    """Excel の範囲分けのとおりに拡散係数を求める（区間1の下限なし、最後の区間の上限を超えたら None）。"""

    n = row['n_ranges']
    if n == 0:
        return 0.0
    if n == 1:
        return math.exp(poly(row['ranges'][0][2], u))
    for i, (lo, hi, co) in enumerate(row['ranges'][:min(n, 3)]):
        if (i == 0 and u <= hi) or (i > 0 and lo < u <= hi):
            return math.exp(poly(co, u))
    return None


def suction_to_rh(p):
    """吸引圧 Pa を相対湿度 % に直す（RR-110 の式(4)）。"""

    return 100.0 * math.exp(-p * _M / (_RHO_W * _R * _T_LAB))


# ---------------------------------------------------------------------------
# 比べた結果のまとめ方
# ---------------------------------------------------------------------------

def ratio_summary(pairs):
    """[(Excel の値, 表の値), ...] から、比 Excel/表 の最小・中央・最大を文字列で返す。"""

    ratios = [e / t for e, t in pairs if t not in (0, None) and e is not None]
    if not ratios:
        return '比べられる点なし'
    return f'{len(ratios)} 点、Excel/表 = {min(ratios):.3g}～{max(ratios):.3g}（中央 {statistics.median(ratios):.3g}）'


def log_error(pairs):
    """比の対数の絶対値の平均（列の対応を決めるときに使う）。"""

    v = [abs(math.log(e / t)) for e, t in pairs if t and e and e > 0 and t > 0]
    return sum(v) / len(v) if v else math.inf


def best_column(row_func, columns: dict):
    """Excel の式に一番近い列の名前を返す。columns は {名前: [(x, 値), ...]}。"""

    scores = {name: log_error([(row_func(x), t) for x, t in points]) for name, points in columns.items()}
    best = min(scores, key=scores.get)
    return best, scores[best]


# ---------------------------------------------------------------------------
# 各物性の照合
# ---------------------------------------------------------------------------

def check_basic(rows):
    print('## 密度・熱伝導率・比熱')
    for no, (rho, lam, src) in T.DENSITY_CONDUCTIVITY.items():
        r = rows[no]
        print(f'No.{no:2d} {r["name"][:40]:40s} 密度 Excel {r["rho"]} / 表 {rho:.1f}（{rho and r["rho"] / rho:.3f}）'
              f'  λ Excel {r["lambda"]} / 表 {lam:.4g}（{r["lambda"] / lam:.3f}）  [{src}]')
    for nos, c, src in T.HEAT_CAPACITY:
        print(f'比熱 No.{nos[0]}～{nos[-1]}: 表 {c}, Excel {sorted({rows[n]["c"] for n in nos})}  [{src}]')
    for no, (name, thickness, mass) in T.MEMBRANE.items():
        print(f'No.{no:2d} メンブレン {name}: 質量/面積 ÷ 厚さ = {mass / thickness:.0f} kg/m3、Excel 密度 {rows[no]["rho"]}')
    k = T.AIR_PERMEABILITY_AERATED / 1.2 * 1.7e-5
    print(f'透気率 No.1: 表 {T.AIR_PERMEABILITY_AERATED} kg/(m·Pa·s) ÷ 空気の密度 1.2 × 粘性係数 1.7e-5 = {k:.3g} m2、'
          f'Excel {rows[1]["k_air"]}')
    print()


def check_vapour(rows):
    print('## 透湿率（湿気伝導率）')
    for no, points in T.VAPOUR_PERMEABILITY.items():
        r = rows[no]
        print(f'No.{no:2d} {r["name"][:40]:40s} {ratio_summary([(vapour(r, rh / 100), v) for rh, v in points])}')
    for no, points in T.VAPOUR_PERMEABILITY_RANGE.items():
        r = rows[no]
        arith = [(vapour(r, rh / 100), (lo + hi) / 2) for rh, lo, hi in points]
        geo = [(vapour(r, rh / 100), math.sqrt(lo * hi)) for rh, lo, hi in points]
        inside = sum(lo <= vapour(r, rh / 100) <= hi for rh, lo, hi in points)
        print(f'No.{no:2d} {r["name"][:40]:40s} 下限と上限の算術平均: {ratio_summary(arith)} / 幾何平均: {ratio_summary(geo)}'
              f' / 下限～上限に入る点 {inside}/{len(points)}')

    # RH 0% の点は除く（Excel の式は x = 0 で 0 になるか計算できないため、表の値と比べられない）
    columns = {name: [(x / 100, v) for x, v in zip(T.FIBREBOARD_RH, vals) if x > 0]
               for name, vals in T.FIBREBOARD_VAPOUR_PERMEABILITY_COLUMNS.items()}
    for no in range(27, 35):
        r = rows[no]
        if r['vapour'][0] == 0:
            print(f'No.{no:2d} {r["name"][:40]:40s} Excel の係数がすべて 0（表 8.2 に Foil Facer はない）')
            continue
        name, err = best_column(lambda x: vapour(r, x), columns)
        pairs = [(vapour(r, x), v) for x, v in columns[name]]
        print(f'No.{no:2d} {r["name"][:40]:40s} 一番近い列 = 表 8.2 の「{name}」: {ratio_summary(pairs)}')

    for no, permeance in T.SIDING_PERMEANCE.items():
        r = rows[no]
        pairs = [(vapour(r, rh / 100), p * T.SIDING_THICKNESS[no]) for rh, p in zip(T.SIDING_RH, permeance)]
        print(f'No.{no:2d} {r["name"][:40]:40s} パーミアンス × 厚さ と比べて: {ratio_summary(pairs)}')

    for no, (name, thickness, _) in T.MEMBRANE.items():
        r = rows[no]
        a, b, c, d = r['vapour']
        if b == 0 and c == 0 and d == 0:
            permeance = a / (thickness / 1000)
            nearest = min(T.MEMBRANE_CONSTANT_PERMEANCE, key=lambda p: abs(math.log(p / permeance)))
            print(f'No.{no:2d} メンブレン {name:4s} 一定値 a ÷ 厚さ = {permeance:.3g} kg/(m2·s·Pa)、表 4.2 の一定の列で一番近い値 {nearest:.3g}')
        else:
            columns = {k: [(rh / 100, p * thickness / 1000) for rh, p in v if rh > 0]
                       for k, v in T.MEMBRANE_VARIABLE_PERMEANCE.items()}
            # RH によらず一定の列も候補に入れる
            for p in T.MEMBRANE_CONSTANT_PERMEANCE:
                columns[f'一定 {p:.3g}'] = [(rh / 100, p * thickness / 1000) for rh in range(10, 101, 10)]
            col, err = best_column(lambda x: vapour(r, x), columns)
            pairs = [(vapour(r, x), v) for x, v in columns[col]]
            print(f'No.{no:2d} メンブレン {name:4s} 一番近い列 = 表 4.2 の列 {col}（パーミアンス × 厚さ）: {ratio_summary(pairs)}')

    r = rows[41]
    print(f'No.41 Gypsum: 表 RH80% のパーミアンス × 厚さ 13 mm = {T.GYPSUM_PERMEANCE_80 * T.GYPSUM_THICKNESS:.3g}、'
          f'Excel a = {r["vapour"][0]}（b～d = {r["vapour"][1:]}）')
    print()


def _isotherm_points(rows_table, nos):
    """表の行（種別, RH または吸引圧, 値の並び）を {No: [(RH%, 含水率, 種別), ...]} に直す。"""

    out = {no: [] for no in nos}
    for kind, x, values in rows_table:
        rh = suction_to_rh(x) if kind == 'P' else x
        for no, v in zip(nos, values):
            out[no].append((rh, v, kind))
    return out


def _isotherm_line(r, points):
    by_kind = {}
    for rh, v, kind in points:
        by_kind.setdefault(kind, []).append((isotherm(r, rh / 100), v))
    return '  '.join(f'{ {"S": "吸湿", "D": "放湿", "P": "加圧板"}[k]}: {ratio_summary(p)}' for k, p in by_kind.items())


def check_isotherm(rows):
    print('## 平衡含水率（加圧板の値は式(4)で 22 °C の相対湿度に直した）')
    data = dict(T.ISOTHERM)
    for no, pts in T.ISOTHERM_PRESSURE_PLATE.items():
        data[no] = data[no] + [(suction_to_rh(p), v, 'P') for p, v in pts]
    data.update(_isotherm_points(T.BRICK_ISOTHERM_ROWS, T.BRICK_NOS))
    data.update(_isotherm_points(T.MORTAR_ISOTHERM_ROWS, T.MORTAR_NOS))
    data.update(_isotherm_points(T.STUCCO_ISOTHERM_ROWS, T.STUCCO_NOS))
    data.update(_isotherm_points(T.SIDING_ISOTHERM_ROWS, T.SIDING_NOS))
    data[40] = [(rh, v, 'S') for rh, v in T.EIFS_ISOTHERM]
    for no in sorted(data):
        r = rows[no]
        print(f'No.{no:2d} {r["name"][:40]:40s} {_isotherm_line(r, data[no])}  （Excel a = {r["iso"][0]}）')

    fb = _isotherm_points(T.FIBREBOARD_ISOTHERM_ROWS, T.FIBREBOARD_ISOTHERM_COLUMNS)
    for no in range(27, 35):
        r = rows[no]
        if r['iso'][0] == 0:
            print(f'No.{no:2d} {r["name"][:40]:40s} Excel の係数がすべて 0（表 8.6 に Foil Facer はない）')
            continue
        columns = {k: [(rh / 100, v) for rh, v, kind in pts if kind != 'P'] for k, pts in fb.items()}
        col, _ = best_column(lambda x: isotherm(r, x), columns)
        print(f'No.{no:2d} {r["name"][:40]:40s} 一番近い列 = 表 8.6 の「{col}」: {_isotherm_line(r, fb[col])}')
    for no in (4, 29, 41, 42):
        print(f'No.{no:2d} {rows[no]["name"][:40]:40s} Excel の平衡含水率の係数 = {rows[no]["iso"]}')
    print()


def _refit(row, points_u):
    """Excel の区間ごとに、ln D を u の多項式で最小二乗近似し直して、Excel の係数と比べる。"""

    lines = []
    n = row['n_ranges']
    ranges = row['ranges'][:max(1, min(n, 3))]
    for i, (lo, hi, co) in enumerate(ranges):
        degree = max((k for k, c in enumerate(co) if c), default=0)
        if n == 1:
            sel = points_u
        else:
            sel = [(u, d) for u, d in points_u if (i == 0 and u <= hi) or (i > 0 and lo < u <= hi)]
        label = f'区間{i + 1}（{lo}～{hi}）' if n > 1 else '全範囲'
        if degree == 0:
            if sel:
                lines.append(f'    {label}: 定数 exp(a) = {math.exp(co[0]):.3g}。この区間の表の値 {sel[0][1]:.3g}～{sel[-1][1]:.3g}')
            else:
                end = min(points_u, key=lambda p: abs(p[0] - (hi if i == 0 else lo)))
                lines.append(f'    {label}: 定数 exp(a) = {math.exp(co[0]):.3g}。区間内に表の点なし（一番近い端の点 u={end[0]:.3g} で {end[1]:.3g}）')
            continue
        if len(sel) <= degree + 1:
            # 点の数が 次数+1 以下だと、どんな係数でも表の点をぴったり通るので、比べる意味がない
            lines.append(f'    {label}: {degree} 次式。区間内の表の点が {len(sel)} 点しかなく、近似し直しても比べられない')
            continue
        u = np.array([p[0] for p in sel])
        y = np.log([p[1] for p in sel])
        fit = np.polynomial.polynomial.polyfit(u, y, degree)
        excel = np.array(co[:degree + 1])
        rms_fit = math.sqrt(np.mean((np.polynomial.polynomial.polyval(u, fit) - y) ** 2))
        rms_excel = math.sqrt(np.mean((np.polynomial.polynomial.polyval(u, excel) - y) ** 2))
        lines.append(f'    {label}: {degree} 次式、表の {len(sel)} 点で近似し直した係数 {np.array2string(fit, precision=5)}'
                     f' / Excel {np.array2string(excel, precision=5)}'
                     f'、ln D の残差 RMS 近似し直し {rms_fit:.3g} / Excel {rms_excel:.3g}')
    return lines


def check_diffusivity(rows):
    print('## 液水拡散係数（含水量 kg/m3 の表は、Excel の乾燥密度で割って kg/kg に直した）')
    data = {}
    for no, pts in T.DIFFUSIVITY_U.items():
        data[no] = pts
    for no, pts in T.DIFFUSIVITY_W_RANGE.items():
        data[no] = [(w / rows[no]['rho'], math.sqrt(lo * hi)) for w, lo, hi in pts]
    for table, nos in ((T.BRICK_DIFFUSIVITY_W, T.BRICK_NOS), (T.MORTAR_DIFFUSIVITY_W, T.MORTAR_NOS),
                       (T.STUCCO_DIFFUSIVITY_W, T.STUCCO_NOS), (T.SIDING_DIFFUSIVITY_W, T.SIDING_DIFFUSIVITY_NOS)):
        for j, no in enumerate(nos):
            data[no] = [(w / rows[no]['rho'], vals[j]) for w, vals in table if vals[j] is not None]

    fb_columns = {name: [(w, vals[j]) for w, vals in T.FIBREBOARD_DIFFUSIVITY_W]
                  for j, name in enumerate(T.FIBREBOARD_DIFFUSIVITY_COLUMNS)}
    for no in range(27, 35):
        r = rows[no]
        if r['n_ranges'] == 0:
            continue
        cols = {k: [(w / r['rho'], d) for w, d in v] for k, v in fb_columns.items()}
        col, _ = best_column(lambda u: diffusivity(r, u) or 1e-300, cols)
        data[no] = cols[col]
        print(f'No.{no:2d} は表 8.7 の「{col}」の列に一番近い（含水量 30～90 kg/m3 の 7 点だけで比べる）')

    for no in sorted(data):
        r = rows[no]
        pts = data[no]
        pairs = [(diffusivity(r, u), d) for u, d in pts]
        print(f'No.{no:2d} {r["name"][:40]:40s} 範囲分け数 {r["n_ranges"]}: {ratio_summary(pairs)}'
              f'（表の含水率 {pts[0][0]:.3g}～{pts[-1][0]:.3g} kg/kg）')
        for line in _refit(r, pts):
            print(line)

    print()
    print('### 範囲分け数 0 の材料の a、および定数の区間の値の重なり')
    constants = {}
    for no, r in rows.items():
        if no > 42:
            continue
        n = r['n_ranges']
        for i, (lo, hi, co) in enumerate(r['ranges'][:max(1, min(n, 3))]):
            if co[0] and not any(co[1:]):
                constants.setdefault(round(co[0], 9), []).append(f'No.{no} 区間{i + 1}' if n else f'No.{no}（範囲分け数 0）')
    for a, where in sorted(constants.items()):
        mark = '  ← 複数の材料で同じ値' if len(where) > 1 else ''
        print(f'  a = {a:.9g}（exp(a) = {math.exp(a):.3g}）: {", ".join(where)}{mark}')
    print(f'  参考: ln(99999) = {math.log(99999):.9g}, ln(1e-16) = {math.log(1e-16):.9g}')
    print()


def main():
    rows = read_excel()
    check_basic(rows)
    check_vapour(rows)
    check_isotherm(rows)
    check_diffusivity(rows)


if __name__ == '__main__':
    main()
