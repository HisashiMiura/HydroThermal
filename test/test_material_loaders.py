import logging
import math
import os

import openpyxl
import pandas as pd
import pytest


from modules.material_loaders import (
    WHAMCAVITY_FILE, load_whamcavity, MEWS_ASHRAE_FILE, load_mews_ashrae, _liquid_diffusivity_from_ranges
)
from modules.material_props import NoLiquidDiffusivity


def test_load_whamcavity():
    """旧 Excel の19材料が、Excel と同じ順で読み込まれることを確認する。"""

    ms = load_whamcavity()

    names = pd.read_excel(WHAMCAVITY_FILE, sheet_name='material_prop')['name'].tolist()

    assert len(ms) == 19
    assert [m.name for m in ms] == names


def test_load_whamcavity_file_not_found(tmp_path):
    """ファイルがないとき、FileNotFoundError になることを確認する。"""

    with pytest.raises(FileNotFoundError):
        load_whamcavity(os.path.join(tmp_path, 'not_found.xlsx'))


def test_load_whamcavity_duplicate_name(tmp_path):
    """名前が重複しているとき、ValueError になることを確認する。

    旧 Excel の先頭の行を末尾にもう一度足した Excel を、一時フォルダ（tmp_path）に作って読む。
    """

    df = pd.read_excel(WHAMCAVITY_FILE, sheet_name='material_prop')
    df = pd.concat([df, df.iloc[[0]]])

    path = os.path.join(tmp_path, 'duplicate.xlsx')
    df.to_excel(path, sheet_name='material_prop', index=False)

    with pytest.raises(ValueError, match='重複'):
        load_whamcavity(path)


def test_liquid_diffusivity_from_ranges_0():
    """範囲分け数 0 のとき、「液水が移動しない」部品になることを確認する。

    No.4 MEWS Building Membrane IA の値（a は「測っていない」の印 ln(99999)）。
    """

    d = _liquid_diffusivity_from_ranges(0, [(None, 0, (math.log(99999), 0, 0, 0, 0, 0, 0))])

    assert isinstance(d, NoLiquidDiffusivity)
    assert d.d_w(0.1) == 0.0


@pytest.mark.parametrize('u', [0.0, 0.05, 0.2])
def test_liquid_diffusivity_from_ranges_1(u):
    """範囲分け数 1 のとき、区間1の式を含水率によらず使うことを確認する。

    No.14 MEWS Brick I - White Concrete Brick の値。Excel の上限は 0 だが使わないので、u > 0 でも式の値になる。
    """

    d = _liquid_diffusivity_from_ranges(1, [(None, 0, (-32.285, 272.18, 0, 0, 0, 0, 0))])

    assert d.d_w(u) == pytest.approx(math.exp(-32.285 + 272.18 * u), rel=1e-12)



# No.1 MEWS Aerated Concrete の値（範囲分け数 3）
_RANGES_NO1 = [
    (None, 0.087, (-18.55994281128587, 0, 0, 0, 0, 0, 0)),
    (0.087, 0.54, (-15.073, -63.312, 324.96, -731.86, 623.3, 0, 0)),
    (0.54, 1.3, (-16.772022118364983, 0, 0, 0, 0, 0, 0)),
]

# No.2 MEWS Oriented Strand Board の値（範囲分け数 2）。区間3 の欄は Excel では 0
_RANGES_NO2 = [
    (None, 0.92, (-18.023, -34.413, 140.32, -285.57, 277.28, -102.33, 0)),
    (0.92, 1.4, (-21.996231512759298, 0, 0, 0, 0, 0, 0)),
    (0, 0, (0, 0, 0, 0, 0, 0, 0)),
]


def _exp_poly(coefficients, u):
    """exp(a + b·u + c·u^2 + d·u^3 + e·u^4 + f·u^5 + g·u^6) を求める（テストの期待値用）。"""
    a, b, c, d, e, f, g = coefficients
    return math.exp(a + b * u + c * u**2 + d * u**3 + e * u**4 + f * u**5 + g * u**6)


@pytest.mark.parametrize('u, i', [
    (0.0, 0),    # 含水率 0 でも区間1（下限は −∞）
    (0.05, 0),
    (0.087, 0),  # 区間1と区間2の境目ちょうどは下側の区間1
    (0.3, 1),
    (0.54, 1),   # 区間2と区間3の境目ちょうどは下側の区間2
    (0.8, 2),
])
def test_liquid_diffusivity_from_ranges_3(u, i):
    """範囲分け数 3 のとき、含水率が入る区間の式を使うことを確認する（No.1 の値）。"""

    d = _liquid_diffusivity_from_ranges(3, _RANGES_NO1)

    assert d.d_w(u) == pytest.approx(_exp_poly(_RANGES_NO1[i][2], u), rel=1e-12)


@pytest.mark.parametrize('u, i', [
    (0.0, 0),
    (0.5, 0),
    (1.2, 1),
])
def test_liquid_diffusivity_from_ranges_2(u, i):
    """範囲分け数 2 のとき、区間1・2の式を使うことを確認する（No.2 の値）。"""

    d = _liquid_diffusivity_from_ranges(2, _RANGES_NO2)

    assert d.d_w(u) == pytest.approx(_exp_poly(_RANGES_NO2[i][2], u), rel=1e-12)


def test_liquid_diffusivity_from_ranges_2_ignores_range3():
    """範囲分け数 2 のとき、区間3 の欄の値は使わないことを確認する（No.2 の値）。

    部品の区間は、区間1・2 と、上限を超える含水率のために足した区間の 3 つになる。
    """

    d = _liquid_diffusivity_from_ranges(2, _RANGES_NO2)

    assert len(d.ranges) == 3
    assert d.ranges[2][:2] == (1.4, math.inf)



# No.76 ASHRAE Gypsum の値（範囲分け数 2）。区間1の上限 0.13 と区間2の下限 0.15 の間にすき間がある
_RANGES_NO76 = [
    (None, 0.13, (-36.841361487904734, 0, 0, 0, 0, 0, 0)),
    (0.15, 0.57, (-30.444, 11.259, 5.7655, -5.3673, 0, 0, 0)),
    (0, 0, (0, 0, 0, 0, 0, 0, 0)),
]


@pytest.mark.parametrize('u, i', [
    (0.12, 0),
    (0.13, 0),   # Excel の区間1の上限
    (0.14, 0),   # すき間の中は、延ばした区間1
    (0.15, 0),   # 延ばした区間1の上限（境目ちょうどは下側の区間）
    (0.16, 1),
])
def test_liquid_diffusivity_from_ranges_gap(u, i):
    """区間と区間のすき間を、前の区間を延ばして埋めることを確認する（No.76 の値）。"""

    d = _liquid_diffusivity_from_ranges(2, _RANGES_NO76)

    assert d.d_w(u) == pytest.approx(_exp_poly(_RANGES_NO76[i][2], u), rel=1e-12)


def test_liquid_diffusivity_from_ranges_gap_log(caplog):
    """すき間を埋めたとき、材料名と値がログに出ることを確認する。すき間がなければ出ない。"""

    with caplog.at_level(logging.INFO, logger='modules.material_loaders'):
        _liquid_diffusivity_from_ranges(2, _RANGES_NO76, name='ASHRAE Gypsum')
        _liquid_diffusivity_from_ranges(3, _RANGES_NO1, name='MEWS Aerated Concrete')

    assert len(caplog.records) == 1
    assert 'ASHRAE Gypsum' in caplog.text
    assert '0.13〜0.15' in caplog.text


def test_liquid_diffusivity_from_ranges_overlap():
    """区間が重なっているとき、ValueError になることを確認する。"""

    ranges = [(None, 0.2, (-20, 0, 0, 0, 0, 0, 0)), (0.1, 0.5, (-18, 0, 0, 0, 0, 0, 0))]

    with pytest.raises(ValueError, match='重なって'):
        _liquid_diffusivity_from_ranges(2, ranges)



# No.49 ASHRAE Fibreboard の値（範囲分け数 2）。最後の区間は 6 次式で、RH 100% の平衡含水率 2.71 が上限 1.5 より大きい
_RANGES_NO49 = [
    (None, 0.058, (-36.841361487904734, 0, 0, 0, 0, 0, 0)),
    (0.094, 1.5, (-21.996, 11.89, -41.679, 72.337, -66.29, 30.704, -5.6596)),
    (0, 0, (0, 0, 0, 0, 0, 0, 0)),
]


@pytest.mark.parametrize('ranges, n, i, upper', [
    (_RANGES_NO1, 3, 2, 1.3),    # 最後の区間が定数
    (_RANGES_NO49, 2, 1, 1.5),   # 最後の区間が多項式
])
@pytest.mark.parametrize('excess', [0.0, 0.1, 1.0, 100.0])
def test_liquid_diffusivity_from_ranges_above_last(ranges, n, i, upper, excess):
    """最後の区間の上限を超える含水率では、上限での値で一定になることを確認する。"""

    d = _liquid_diffusivity_from_ranges(n, ranges)

    assert d.d_w(upper + excess) == pytest.approx(_exp_poly(ranges[i][2], upper), rel=1e-12)


def test_liquid_diffusivity_from_ranges_above_last_log(caplog):
    """RH 100% の平衡含水率が最後の区間の上限より大きい材料だけ、ログが出ることを確認する。"""

    with caplog.at_level(logging.INFO, logger='modules.material_loaders'):
        _liquid_diffusivity_from_ranges(2, _RANGES_NO49, name='ASHRAE Fibreboard', u_saturation=2.71)
        _liquid_diffusivity_from_ranges(3, _RANGES_NO1, name='MEWS Aerated Concrete', u_saturation=0.787)

    messages = [r.getMessage() for r in caplog.records if '最後の区間' in r.getMessage()]
    assert len(messages) == 1
    assert 'ASHRAE Fibreboard' in messages[0]



# No.27 MEWS Fibreboard I - Paper Facer の値（範囲分け数 4。値は3区間分しかない）
_RANGES_NO27 = [
    (None, 0.11, (-19.890356714011308, 0, 0, 0, 0, 0, 0)),
    (0.11, 0.32, (5.6968, -677.38, 6968.5, -35046, 85320, -79720, 0)),
    (0.32, 0.53, (124.47, -1250, 4080.2, -5961, 3265.1, 0, 0)),
]

def test_liquid_diffusivity_from_ranges_4(caplog):
    """範囲分け数 4 のとき、3 として読み、ログ（WARNING）が出ることを確認する（No.27 の値）。"""

    with caplog.at_level(logging.WARNING, logger='modules.material_loaders'):
        d4 = _liquid_diffusivity_from_ranges(4, _RANGES_NO27, name='MEWS Fibreboard I - Paper Facer')

    d3 = _liquid_diffusivity_from_ranges(3, _RANGES_NO27)

    assert d4 == d3
    assert 'MEWS Fibreboard I - Paper Facer' in caplog.text
    # 区間3の上限 0.53 を超えたら、0.53 での値で一定
    assert d4.d_w(0.6) == pytest.approx(_exp_poly(_RANGES_NO27[2][2], 0.53), rel=1e-12)


@pytest.mark.parametrize('n_ranges', [-1, 5])
def test_liquid_diffusivity_from_ranges_invalid(n_ranges):
    """範囲分け数が 0〜4 でないとき、ValueError になることを確認する。"""

    with pytest.raises(ValueError, match='範囲分け数'):
        _liquid_diffusivity_from_ranges(n_ranges, _RANGES_NO1)


# ---------------------------------------------------------------------------------------------------------------------
# 新しい Excel（materials_mews_ashrae.xlsx）のローダー
# ---------------------------------------------------------------------------------------------------------------------

# Excel が計算した値を読むため data_only=True で開く
_wb_mews_ashrae = openpyxl.load_workbook(MEWS_ASHRAE_FILE, data_only=True)

# 「データファイル用」シートの {No: 行の値}（3行目から）
_data_rows = {
    row[0]: row
    for row in _wb_mews_ashrae['データファイル用'].iter_rows(min_row=3, max_col=44, values_only=True)
    if row[0] is not None
}

# 読み込んだ材料（No の順に並ぶ）
_mews_ashrae_materials = dict(zip(_data_rows.keys(), load_mews_ashrae()))

# No.80 は No.62 と同じ名前なので、ローダーが名前の後ろに密度を足す
_NAME_NO80 = 'ASHRAE Low-density Glassfibre Batt Insulation (32 kg/m3)'


def _read_sheet_values(sheet_name: str, first_col: int):
    """「湿気伝導率」「平衡含水率」シートから、相対湿度と Excel が計算した値を読む。

    2行目に相対湿度（0.0～1.0）、3行目から A列 No と first_col 列以降の計算値が並ぶ。
    No が数値でない行より下はグラフ表示用の別の表なので読まない。

    Returns:
        相対湿度のリスト、{No: [各相対湿度での値]}
    """

    ws = _wb_mews_ashrae[sheet_name]
    xs = [ws.cell(2, col).value for col in range(first_col, ws.max_column + 1)]
    expected = {}
    for row in range(3, ws.max_row + 1):
        no = ws.cell(row, 1).value
        if not isinstance(no, int):
            break
        expected[no] = [ws.cell(row, col).value for col in range(first_col, ws.max_column + 1)]

    return xs, expected


_vapor_xs, _vapor_expected = _read_sheet_values('湿気伝導率', first_col=7)
_isotherm_xs, _isotherm_expected = _read_sheet_values('平衡含水率', first_col=13)


def test_load_mews_ashrae():
    """81材料が Excel と同じ順で読み込まれ、No.80 だけ名前に密度が足されることを確認する。"""

    names = [row[1] for row in _data_rows.values()]
    names[80] = _NAME_NO80

    assert len(_mews_ashrae_materials) == 81
    assert [m.name for m in _mews_ashrae_materials.values()] == names


@pytest.mark.parametrize('no', list(_data_rows.keys()))
def test_load_mews_ashrae_basic_values(no):
    """密度・比熱・熱伝導率が Excel の値と同じであることを確認する。"""

    m = _mews_ashrae_materials[no]
    row = _data_rows[no]

    assert m.rho == row[2]
    assert m.c == row[3]
    assert m.lambda_h == row[4]


@pytest.mark.parametrize('no', list(_vapor_expected.keys()))
def test_load_mews_ashrae_vapor_conductivity(no):
    """読み込んだ材料の湿気伝導率が、Excel の「湿気伝導率」シートの計算値と同じであることを確認する。"""

    m = _mews_ashrae_materials[no]

    for x, expected in zip(_vapor_xs, _vapor_expected[no]):

        # d が負の材料（No.28）は x = 0 で Excel でも計算できない（部品のテストで確認済み）
        if expected == '#DIV/0!':
            continue

        assert m.get_lambda_dsh_m(x * 100) == pytest.approx(expected, rel=1e-9, abs=0), f'x={x}'


@pytest.mark.parametrize('no', list(_isotherm_expected.keys()))
def test_load_mews_ashrae_isotherm(no):
    """読み込んだ材料の平衡含水率が、Excel の「平衡含水率」シートの計算値と同じであることを確認する。"""

    m = _mews_ashrae_materials[no]

    for x, expected in zip(_isotherm_xs, _isotherm_expected[no]):
        assert m.get_u(x * 100) == pytest.approx(expected, rel=1e-9, abs=0), f'x={x}'


def _read_diffusivity_sheet():
    """「拡散係数」シートから、範囲分け数と区間ごとの範囲・係数を読む（Excel が「データファイル用」から写した値）。

    4行目から A列 No、C列 範囲分け数、D列 区間1の範囲(大)、E～K列 係数、L・M列 区間2の範囲、N～T列 係数、
    U・V列 区間3の範囲、W～AC列 係数が並ぶ（No.1～79。No.80 はこのシートにない）。

    Returns:
        {No: (範囲分け数, [(下限, 上限, 係数), ...])}
    """

    ws = _wb_mews_ashrae['拡散係数']
    result = {}
    for row in ws.iter_rows(min_row=4, max_col=29, values_only=True):
        if not isinstance(row[0], int):
            break
        ranges = [(0.0, row[3], row[4:11]), (row[11], row[12], row[13:20]), (row[20], row[21], row[22:29])]
        result[row[0]] = (row[2], ranges)

    return result


_diffusivity_sheet = _read_diffusivity_sheet()


@pytest.mark.parametrize('no', list(_diffusivity_sheet.keys()))
def test_load_mews_ashrae_liquid_diffusivity(no):
    """読み込んだ材料の液水拡散係数が、「拡散係数」シートの範囲と係数から求めた値と同じであることを確認する。

    区間ごとに、範囲の真ん中の含水率で比べる（区間1 は 0 から上限までの真ん中）。
    範囲分け数 0 は 0、範囲分け数 1 は区間1の式を含水率 0.1 で比べる。範囲分け数 4（No.27）は 3 区間分を比べる。
    """

    m = _mews_ashrae_materials[no]
    n_ranges, ranges = _diffusivity_sheet[no]

    if n_ranges == 0:
        assert m.liquid_diffusivity.d_w(0.1) == 0.0
        return

    if n_ranges == 1:
        points = [(0.1, ranges[0][2])]
    else:
        points = [((lower + upper) / 2, coefficients) for lower, upper, coefficients in ranges[:min(n_ranges, 3)]]

    for u, coefficients in points:
        expected = math.exp(sum(c * u**k for k, c in enumerate(coefficients)))
        assert m.liquid_diffusivity.d_w(u) == pytest.approx(expected, rel=1e-12), f'u={u}'


def test_load_mews_ashrae_warnings(caplog):
    """No.27（範囲分け数 4）と No.80（同じ名前）で、ログ（WARNING）が出ることを確認する。"""

    with caplog.at_level(logging.WARNING, logger='modules.material_loaders'):
        load_mews_ashrae()

    messages = [r.getMessage() for r in caplog.records if r.levelno == logging.WARNING]

    assert len(messages) == 2
    assert 'No.27' in messages[0]
    assert 'No.80' in messages[1] and _NAME_NO80 in messages[1]


def test_load_mews_ashrae_file_not_found(tmp_path):
    """ファイルがないとき、FileNotFoundError になることを確認する。"""

    with pytest.raises(FileNotFoundError):
        load_mews_ashrae(os.path.join(tmp_path, 'not_found.xlsx'))


def test_load_mews_ashrae_duplicate_name(tmp_path):
    """名前に密度を足しても重なるとき、ValueError になることを確認する。

    「データファイル用」シートの末尾に No.80 の行をもう一度足した Excel を、一時フォルダ（tmp_path）に作って読む。
    3つ目の同じ名前に密度を足すと、No.80 に足した名前と重なる。
    """

    df = pd.read_excel(MEWS_ASHRAE_FILE, sheet_name='データファイル用', header=None)
    df = pd.concat([df, df[df[0] == 80]])

    path = os.path.join(tmp_path, 'duplicate.xlsx')
    df.to_excel(path, sheet_name='データファイル用', header=False, index=False)

    with pytest.raises(ValueError, match='重複'):
        load_mews_ashrae(path)
